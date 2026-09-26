"""Tests for filling the Word template: placeholders, delivery date, line item table."""

from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.table import Table

from src.backend.docgen import (
    fill_table_with_line_items,
    replace_delivery_date,
    replace_placeholders,
)
from src.backend.models import LineItem
from src.backend.placeholders import (
    PH_DATE_TODAY,
    PH_DELIVERY_DATE,
    PH_DOCTYPE,
    PH_PROJECT_NO,
    PH_RECEIPT_NO,
    PH_SUM_GROSS,
    PH_SUM_NET,
)

FIXTURES_DIR = Path(__file__).parents[1] / "fixtures"
TEMPLATE_PATH = FIXTURES_DIR / "Vordruck_sample.docx"


@pytest.fixture
def sample_template() -> DocxDocument:
    return Document(str(TEMPLATE_PATH))


def paragraph_texts(doc: DocxDocument) -> list[str]:
    return [paragraph.text for paragraph in doc.paragraphs]


def address_table(doc: DocxDocument) -> Table:
    return doc.tables[1]


def item_table(doc: DocxDocument) -> Table:
    return doc.tables[2]


def row_texts(table: Table, row_index: int) -> list[str]:
    return [cell.text for cell in table.rows[row_index].cells]


def hours_item(
    order_number: str,
    quantity: str,
    description: str,
    unit_price: str,
    total_price: str,
) -> LineItem:
    return LineItem(
        kind="hours",
        order_number=order_number,
        quantity=Decimal(quantity),
        description=description,
        unit_price=Decimal(unit_price),
        total_price=Decimal(total_price),
    )


def material_item(order_number: str, description: str, amount: str) -> LineItem:
    return LineItem(
        kind="material",
        order_number=order_number,
        quantity=Decimal(1),
        description=description,
        unit_price=Decimal(amount),
        total_price=Decimal(amount),
    )


MASTER_HOURS = hours_item(
    "90010004", "2.50", "Meisterstunde zu Auftrag Nr. 90010004", "84.90", "212.25"
)
HELPER_HOURS = hours_item(
    "111", "3.00", "Helferstunde zu Auftrag Nr. 111", "48.00", "144.00"
)
MATERIAL = material_item("111", "Material zu Auftrag Nr. 111", "270.00")


# — Placeholders ——————————————————————————————————————————————————————————————


def test_placeholder_in_body_paragraph_is_replaced(sample_template):
    replace_placeholders(sample_template, {PH_DATE_TODAY: "26.09.2026"})

    assert "Datum:\t26.09.2026" in paragraph_texts(sample_template)


def test_placeholder_in_table_cell_is_replaced(sample_template):
    replace_placeholders(sample_template, {PH_SUM_GROSS: "303,09€"})

    assert item_table(sample_template).cell(4, 4).text == "303,09€"


def test_all_placeholders_in_one_run_are_replaced(sample_template):
    replace_placeholders(
        sample_template,
        {PH_DOCTYPE: "Rechnung", PH_PROJECT_NO: "4711", PH_RECEIPT_NO: "4000012345"},
    )

    assert "Rechnung Nr. 4711 – 4000012345" in paragraph_texts(sample_template)
    assert (
        "Auftragsnummer:\t4711 / 4000012345"
        in address_table(sample_template).cell(1, 2).text
    )


def test_placeholder_is_replaced_at_every_occurrence(sample_template):
    replace_placeholders(sample_template, {PH_SUM_NET: "254,70€"})

    assert item_table(sample_template).cell(2, 4).text == "254,70€"
    assert item_table(sample_template).cell(3, 0).text == "Ust. 19% auf 254,70€ netto"


# — Delivery date —————————————————————————————————————————————————————————————


def test_delivery_date_split_across_runs_is_replaced(sample_template):
    replace_delivery_date(sample_template, {PH_DELIVERY_DATE: "10.10.2026"})

    assert "Liefertermin: 10.10.2026" in paragraph_texts(sample_template)


def test_text_around_split_delivery_date_is_kept():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Liefertermin: <Liefer")
    paragraph.add_run("datum> ab Werk")

    replace_delivery_date(doc, {PH_DELIVERY_DATE: "10.10.2026"})

    assert paragraph.text == "Liefertermin: 10.10.2026 ab Werk"


# — Line item table ———————————————————————————————————————————————————————————


def test_single_item_fills_the_template_row(sample_template):
    fill_table_with_line_items(sample_template, [MASTER_HOURS])

    table = item_table(sample_template)
    assert row_texts(table, 1) == [
        "1",
        "2,5",
        "Meisterstunde zu Auftrag Nr. 90010004",
        "84,90€",
        "212,25€",
    ]
    assert len(table.rows) == 5


def test_items_are_inserted_in_order_before_the_sum_rows(sample_template):
    fill_table_with_line_items(sample_template, [MASTER_HOURS, HELPER_HOURS, MATERIAL])

    table = item_table(sample_template)
    assert row_texts(table, 1) == [
        "1",
        "2,5",
        "Meisterstunde zu Auftrag Nr. 90010004",
        "84,90€",
        "212,25€",
    ]
    assert row_texts(table, 2) == [
        "2",
        "3",
        "Helferstunde zu Auftrag Nr. 111",
        "48,00€",
        "144,00€",
    ]
    assert row_texts(table, 3) == [
        "3",
        "1",
        "Material zu Auftrag Nr. 111",
        "270,00€",
        "270,00€",
    ]
    assert row_texts(table, 4)[0] == "Summe"
    assert len(table.rows) == 7


def test_price_columns_are_right_aligned(sample_template):
    fill_table_with_line_items(sample_template, [MASTER_HOURS, HELPER_HOURS])

    for row in item_table(sample_template).rows[1:3]:
        for cell in row.cells[3:]:
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_item_cells_are_set_in_calibri_9_bold(sample_template):
    fill_table_with_line_items(sample_template, [MASTER_HOURS, HELPER_HOURS])

    for row in item_table(sample_template).rows[1:3]:
        for cell in row.cells:
            run = cell.paragraphs[0].runs[0]
            assert run.font.name == "Calibri"
            assert run.font.size == Pt(9)
            assert run.font.bold is True


def test_template_without_item_table_raises():
    doc = Document()

    with pytest.raises(ValueError):
        fill_table_with_line_items(doc, [MASTER_HOURS])
