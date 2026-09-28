"""End-to-end tests: sample CSV in, finished Word documents out."""

import shutil
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table, _Cell

from src.backend import docgen, paths, services
from src.backend.config import Config
from src.backend.input_args import (
    create_delivery_args,
    create_invoice_args,
    create_offer_args,
)
from src.backend.services import (
    generate_delivery,
    generate_invoice_and_order,
    generate_offer,
)

FIXTURES_DIR = Path(__file__).parents[1] / "fixtures"
PROJECT_FOLDER = "1234 - Schutzgitter Mischer 3"

# CSV row 11 has a blank order number and is skipped
ITEM_ROWS = [
    ["1", "1", "Meisterstunde zu Auftrag Nr. 123", "84,90€", "84,90€"],
    ["2", "1", "Material zu Auftrag Nr. 123", "95,00€", "95,00€"],
    ["3", "11", "Meisterstunde zu Auftrag Nr. 90010001", "84,90€", "933,90€"],
    ["4", "5", "Meisterstunde zu Auftrag Nr. 90010002", "84,90€", "424,50€"],
    ["5", "7", "Meisterstunde zu Auftrag Nr. 90010003", "84,90€", "594,30€"],
    ["6", "4", "Meisterstunde Sonntag zu Auftrag Nr. 90010003", "139,90€", "559,60€"],
    ["7", "7", "Helferstunde zu Auftrag Nr. 111", "48,00€", "336,00€"],
    ["8", "1", "Material zu Auftrag Nr. 111", "30,00€", "30,00€"],
    ["9", "3", "Helferstunde zu Auftrag Nr. 111", "48,00€", "144,00€"],
    ["10", "1", "Material zu Auftrag Nr. 111", "270,00€", "270,00€"],
    ["11", "2,5", "Meisterstunde zu Auftrag Nr. 90010004", "84,90€", "212,25€"],
    ["12", "2", "Meisterstunde zu Auftrag Nr. 90010005", "84,90€", "169,80€"],
    ["13", "1", "Material zu Auftrag Nr. 90010006", "145,00€", "145,00€"],
]

SUM_ROWS = [
    ["Summe", "Summe", "", "", "3.999,25€"],
    [*["Ust. 19% auf 3.999,25€ netto"] * 3, "", "759,86€"],
    [*["Gesamtbetrag"] * 3, "", "4.759,11€"],
]


@pytest.fixture(autouse=True)
def patch_template_paths(tmp_path, monkeypatch):
    """The template paths are fixed in the code until the path resolution is rebuilt."""
    monkeypatch.setattr(docgen, "VORDRUCK_PATH", FIXTURES_DIR / "Vordruck_sample.docx")
    monkeypatch.setattr(paths, "INTERMEDIATE_ROOT", tmp_path / "intermediate")


@pytest.fixture(autouse=True)
def no_pdf(monkeypatch):
    """PDF rendering is checked by hand."""
    monkeypatch.setattr(services, "render_pdf", lambda docx_path, messages: None)


@pytest.fixture
def sample_data_root(tmp_path) -> Path:
    """Data root in `tmp_path` with one project folder holding the sample CSV."""
    data_root = tmp_path / "Daten"
    project_dir = data_root / PROJECT_FOLDER
    project_dir.mkdir(parents=True)
    shutil.copy(FIXTURES_DIR / "heinrich_zeiterfassung_sample.csv", project_dir)
    return data_root


@pytest.fixture
def sample_config_with_data_root(sample_config, sample_data_root) -> Config:
    return replace(sample_config, data_root=sample_data_root)


def open_document(data_root: Path, filename: str) -> DocxDocument:
    return Document(str(data_root / PROJECT_FOLDER / filename))


def info_block(doc: DocxDocument) -> _Cell:
    return doc.tables[1].cell(1, 2)


def item_table(doc: DocxDocument) -> Table:
    return doc.tables[2]


def item_table_rows(doc: DocxDocument) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in item_table(doc).rows[1:]]


def texts(doc: DocxDocument) -> list[str]:
    """Paragraph texts of the body and the info block."""
    return [p.text for p in doc.paragraphs + info_block(doc).paragraphs]


def assert_items_and_sums(doc: DocxDocument) -> None:
    assert item_table_rows(doc) == ITEM_ROWS + SUM_ROWS


# — Angebot ———————————————————————————————————————————————————————————————————


def test_offer(sample_config_with_data_root, sample_data_root):
    generate_offer(
        create_offer_args("1234"), date(2025, 8, 1), sample_config_with_data_root
    )

    doc = open_document(sample_data_root, "Angebot Nr. 1234.docx")
    doc_texts = texts(doc)
    assert_items_and_sums(doc)
    assert "Datum:\t01.08.2025" in doc_texts
    assert "Angebot Nr. 1234 – " in doc_texts
    assert "Auftragsnummer:\t1234 / " in doc_texts
    assert (
        "Vielen Dank für ihr Vertrauen in Heinrich Metallbau. "
        "Wir bieten folgende Positionen an: "
    ) in doc_texts
    assert "Liefertermin: 15.08.2025" in doc_texts


# — Lieferschein ——————————————————————————————————————————————————————————————


def test_delivery_note(sample_config_with_data_root, sample_data_root):
    args = create_delivery_args("1234", "4100012345")
    generate_delivery(args, date(2025, 8, 15), sample_config_with_data_root)

    doc = open_document(sample_data_root, "Lieferschein Nr. 1234.docx")
    doc_texts = texts(doc)
    assert_items_and_sums(doc)
    assert "Datum:\t15.08.2025" in doc_texts
    assert "Lieferschein Nr. 1234 – 4100012345" in doc_texts
    assert "Auftragsnummer:\t1234 / 4100012345" in doc_texts
    assert (
        "Vielen Dank für ihr Vertrauen in Heinrich Metallbau. "
        "Wir liefern folgende Positionen an: "
    ) in doc_texts
    assert "Liefertermin: 15.08.2025" in doc_texts


# — Rechnung und Auftragsbestätigung ——————————————————————————————————————————


@pytest.mark.parametrize(
    ("filename", "doctype", "header"),
    [
        (
            "Rechnung Nr. 1234 - 4100012345.docx",
            "Rechnung",
            "Wir bitten um Ausgleich der folgenden Positionen:",
        ),
        (
            "Auftragsbestätigung Nr. 1234 - 4100012345.docx",
            "Auftragsbestätigung",
            "Wir bestätigen den Auftrag über folgende Positionen:",
        ),
    ],
)
def test_invoice_and_order_confirmation_after_offer(
    sample_config_with_data_root, sample_data_root, filename, doctype, header
):
    """The delivery date is inherited from the offer via the intermediate template."""
    generate_offer(
        create_offer_args("1234"), date(2025, 8, 1), sample_config_with_data_root
    )
    args = create_invoice_args("1234", "4100012345")
    generate_invoice_and_order(args, date(2025, 8, 20), sample_config_with_data_root)

    doc = open_document(sample_data_root, filename)
    doc_texts = texts(doc)
    assert_items_and_sums(doc)
    assert "Datum:\t20.08.2025" in doc_texts
    assert f"{doctype} Nr. 1234 – 4100012345" in doc_texts
    assert "Auftragsnummer:\t1234 / 4100012345" in doc_texts
    assert (
        f"Vielen Dank für ihr Vertrauen in Heinrich Metallbau. {header} "
    ) in doc_texts
    assert "Liefertermin: 15.08.2025" in doc_texts
